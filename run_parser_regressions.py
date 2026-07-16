import ast
import getpass
import io
import json
import os
import re
import sys
from copy import copy
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


APP_PATH = Path(__file__).with_name("app0518_smart_local.py")
EXAMPLES_PATH = Path(__file__).with_name("parser_regression_examples.md")


class FakeStreamlit:
    def cache_data(self, *args, **kwargs):
        def decorator(fn):
            return fn

        return decorator


def load_parser_namespace():
    source = APP_PATH.read_text()
    tree = ast.parse(source)
    keep_assignments = {
        "DEFAULT_PRODUCT_CATALOG_PATH",
        "DEFAULT_TEMPLATE_PATH",
        "PRODUCT_HEADER_ALIASES",
        "PRODUCT_GROUP_ALIASES",
        "FLAVOR_TEXT_ALIASES",
        "ALL_IN_STOCK_SERIES_PREFIX",
        "SHIPMENT_INCREMENT",
    }
    selected_nodes = []
    for node in tree.body:
        if isinstance(node, ast.Assign):
            names = {target.id for target in node.targets if isinstance(target, ast.Name)}
            if names & keep_assignments:
                selected_nodes.append(node)
        elif isinstance(node, ast.FunctionDef):
            selected_nodes.append(node)

    module = ast.Module(body=selected_nodes, type_ignores=[])
    ast.fix_missing_locations(module)
    namespace = {
        "__file__": str(APP_PATH),
        "st": FakeStreamlit(),
        "pd": pd,
        "re": re,
        "json": json,
        "os": os,
        "sys": sys,
        "io": io,
        "getpass": getpass,
        "Path": Path,
        "SequenceMatcher": SequenceMatcher,
        "copy": copy,
        "datetime": datetime,
        "load_workbook": load_workbook,
    }
    exec(compile(module, str(APP_PATH), "exec"), namespace)
    return namespace


def iter_examples(markdown_text):
    parts = re.split(r"^## ", markdown_text, flags=re.MULTILINE)
    for part in parts[1:]:
        title = part.splitlines()[0].strip()
        blocks = re.findall(r"```text\n(.*?)\n```", part, flags=re.DOTALL)
        if len(blocks) < 2:
            continue
        raw_text = blocks[0]
        expected_skus = [
            line.strip()
            for line in blocks[1].splitlines()
            if line.strip()
        ]
        expected_review_inputs = []
        if "Expected review candidates" in part and len(blocks) >= 3:
            expected_review_inputs = [
                line.strip()
                for line in blocks[2].splitlines()
                if line.strip()
            ]
        yield title, raw_text, expected_skus, expected_review_inputs


def main():
    namespace = load_parser_namespace()
    catalog_df, catalog_error = namespace["load_product_catalog"](
        str(namespace["DEFAULT_PRODUCT_CATALOG_PATH"])
    )
    if catalog_error:
        raise RuntimeError(catalog_error)

    failures = []
    for title, raw_text, expected_skus, expected_review_inputs in iter_examples(
        EXAMPLES_PATH.read_text()
    ):
        rows, parser_name, review_rows = namespace["parse_demand_with_local_fallback"](
            raw_text,
            catalog_df,
        )
        found_skus = {str(row.get("SKU", "")).strip() for row in rows}
        missing = [sku for sku in expected_skus if sku not in found_skus]
        needed_review_rows = [
            row for row in review_rows if bool(row.get("Need Review", False))
        ]
        missing_review_inputs = []
        for expected_input in expected_review_inputs:
            expected_norm = namespace["normalize_match_text"](expected_input)
            matched = any(
                expected_norm
                and expected_norm in namespace["normalize_match_text"](
                    " ".join([
                        str(row.get("Input", "")),
                        str(row.get("Parsed Flavor", "")),
                        str(row.get("Matched Flavor", "")),
                    ])
                )
                for row in needed_review_rows
            )
            if not matched:
                missing_review_inputs.append(expected_input)
        total_quantity = sum(float(row.get("Quantity", 0) or 0) for row in rows)
        review_count = len(needed_review_rows)
        status = "PASS" if not missing and not missing_review_inputs else "FAIL"
        print(
            f"{status}: {title} | parser={parser_name} | "
            f"rows={len(rows)} | qty={total_quantity:g} | review={review_count}"
        )
        if missing:
            print(f"  Missing expected SKU(s): {', '.join(missing)}")
        if missing_review_inputs:
            print(
                "  Missing expected review candidate(s): "
                + ", ".join(missing_review_inputs)
            )
        if missing or missing_review_inputs:
            failures.append((title, missing))

    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
